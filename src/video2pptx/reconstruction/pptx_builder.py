from __future__ import annotations

from pathlib import Path
from tempfile import NamedTemporaryFile

from video2pptx.layout.geometry import pixel_bbox_to_inches, place_logo


THEMES = {
    "original": ("FFFFFF", "111827", "2563EB"),
    "corporate": ("FFFFFF", "1F2937", "2F75B5"),
    "modern-minimal": ("FAFAFA", "111827", "10B981"),
    "premium-dark": ("111827", "F9FAFB", "D4AF37"),
    "custom-brand": None,
    "hybrid": ("FFFFFF", "1F2937", "2F75B5"),
}


def _rgb(value):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(value.lstrip("#"))


def _set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _fit_image(slide, path, box):
    from PIL import Image
    from pptx.util import Inches
    x, y, w, h = box
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio >= box_ratio:
        ph, pw = w / ratio, w
        px, py = x, y + (h - ph) / 2
    else:
        pw, ph = h * ratio, h
        px, py = x + (w - pw) / 2, y
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def _logo_asset(path, opacity):
    if not path:
        return None, None
    path = Path(path)
    if path.suffix.lower() == ".svg":
        try:
            import cairosvg
            tmp = NamedTemporaryFile(suffix=".png", delete=False)
            tmp.close()
            cairosvg.svg2png(url=str(path), write_to=tmp.name)
            path = Path(tmp.name)
        except (ImportError, OSError) as exc:
            return None, f"SVG renderer unavailable ({type(exc).__name__}); logo was skipped. Use PNG or install CairoSVG with its native Cairo library."
    if opacity < 0.999:
        from PIL import Image
        image = Image.open(path).convert("RGBA")
        alpha = image.getchannel("A").point(lambda p: int(p * opacity))
        image.putalpha(alpha)
        tmp = NamedTemporaryFile(suffix=".png", delete=False)
        tmp.close(); image.save(tmp.name); path = Path(tmp.name)
    return path, None


def _add_branding(slide, slide_index, slide_count, logo, brand, slide_size, occupied):
    from PIL import Image
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    sw, sh = slide_size
    if logo:
        with Image.open(logo) as image:
            aspect = image.width / image.height
        placement = place_logo(brand.logo_position, slide_size, aspect, brand.logo_scale,
                               brand.logo_margin_inches, brand.sidebar_width_inches, occupied)
        if placement.sidebar:
            x = sw - brand.sidebar_width_inches if placement.sidebar == "right" else 0
            rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), 0,
                                          Inches(brand.sidebar_width_inches), Inches(sh))
            rect.fill.solid(); rect.fill.fore_color.rgb = _rgb(brand.primary_color); rect.line.fill.background()
        slide.shapes.add_picture(str(logo), Inches(placement.x), Inches(placement.y),
                                 Inches(placement.width), Inches(placement.height))
    footer = brand.footer_text.strip()
    if footer or brand.show_slide_numbers:
        box = slide.shapes.add_textbox(Inches(0.35), Inches(sh - 0.3), Inches(sw - 0.7), Inches(0.2))
        p = box.text_frame.paragraphs[0]
        p.text = footer + (("  ·  " if footer else "") + str(slide_index + 1) if brand.show_slide_numbers else "")
        p.font.name = brand.font_family; p.font.size = Pt(8); p.font.color.rgb = _rgb(brand.text_color)
        p.alignment = PP_ALIGN.RIGHT


def _add_editable_text(slide, element, source_size, target_box, brand, theme_colors):
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Inches, Pt
    sx, sy, sw, sh = target_box
    x, y, w, h = pixel_bbox_to_inches(element.bbox, source_size, (sw, sh))
    box = slide.shapes.add_textbox(Inches(sx + x), Inches(sy + y), Inches(max(w, 0.25)), Inches(max(h, 0.18)))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]; p.text = element.text or ""; p.font.name = brand.font_family
    p.font.size = Pt(max(8, min(28, h * 36))); p.font.color.rgb = _rgb(theme_colors[1])
    p.alignment = PP_ALIGN.RIGHT if element.rtl else PP_ALIGN.LEFT
    return (sx + x, sy + y, w, h)


def build_presentation(model, output_path, kind, brand, logo_path=None, theme="corporate",
                       preserve_animation_stages=False):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    prs = Presentation()
    ratio = model.width_px / model.height_px
    slide_w = 13.333 if ratio >= 1.5 else 10.0
    slide_h = slide_w / ratio
    prs.slide_width, prs.slide_height = Inches(slide_w), Inches(slide_h)
    colors = (brand.background_color, brand.text_color, brand.accent_color) if theme == "custom-brand" else THEMES.get(theme, THEMES["corporate"])
    logo, logo_warning = _logo_asset(logo_path, brand.logo_opacity)
    for index, source in enumerate(model.slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_background(slide, colors[0])
        exclude_logo = (index == 0 and brand.exclude_title_slide) or (index == len(model.slides) - 1 and brand.exclude_closing_slide)
        sidebar = brand.sidebar_width_inches if logo and brand.logo_position in {"right-sidebar", "left-sidebar"} and not exclude_logo else 0
        content_x = sidebar if brand.logo_position == "left-sidebar" else 0
        content_w = slide_w - sidebar
        content_h = slide_h - 0.34
        occupied = []
        if kind in {"reference", "hybrid"} or not any(e.editable for e in source.elements):
            pad = 0 if kind == "reference" else 0.18
            pic_box = (content_x + pad, pad, content_w - 2 * pad, content_h - pad)
            _fit_image(slide, source.reference_frame, pic_box)
            occupied.append(pic_box)
        else:
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(content_x), 0, Inches(0.08), Inches(content_h))
            accent.fill.solid(); accent.fill.fore_color.rgb = _rgb(colors[2]); accent.line.fill.background()
            target = (content_x + 0.35, 0.25, content_w - 0.7, content_h - 0.45)
            for element in source.elements:
                if element.kind == "text" and element.editable and element.text:
                    occupied.append(_add_editable_text(slide, element, (source.width_px, source.height_px), target, brand, colors))
        if kind != "reference" and not exclude_logo:
            _add_branding(slide, index, len(model.slides), logo, brand, (slide_w, slide_h), occupied)
        elif kind != "reference":
            _add_branding(slide, index, len(model.slides), None, brand, (slide_w, slide_h), occupied)
        if logo_warning:
            source.warnings.append(logo_warning)
        notes = slide.notes_slide.notes_text_frame
        notes.text = "Reconstruction metadata: " + "; ".join(source.warnings + [f"Animations detected: {len(source.animations)}"])
    output_path = Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path

