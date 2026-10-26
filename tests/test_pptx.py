from pathlib import Path
import zipfile

from PIL import Image
from pptx import Presentation

from video2pptx.branding import BrandConfig
from video2pptx.models import Element, PresentationModel, Slide
from video2pptx.reconstruction import build_presentation


def test_pptx_file_validity_and_editable_text(tmp_path):
    image = tmp_path / "slide.jpg"; Image.new("RGB", (1600, 900), "white").save(image)
    element = Element("text", (100, 100, 600, 80), "Editable title", .99, "eng", "title", True)
    slide = Slide(0, 0, 1, str(image), 1600, 900, elements=[element])
    model = PresentationModel("sample.mp4", "abc", 1600, 900, 30, 1, [slide])
    output = tmp_path / "result.pptx"
    build_presentation(model, output, "branded", BrandConfig(), theme="corporate")
    assert zipfile.is_zipfile(output)
    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    assert any(shape.has_text_frame and "Editable title" in shape.text for shape in presentation.slides[0].shapes)


def test_source_ratio_preserved(tmp_path):
    image = tmp_path / "slide.jpg"; Image.new("RGB", (1200, 900), "white").save(image)
    model = PresentationModel("x.mp4", "x", 1200, 900, 30, 1, [Slide(0, 0, 1, str(image), 1200, 900)])
    output = tmp_path / "ratio.pptx"; build_presentation(model, output, "reference", BrandConfig())
    prs = Presentation(output)
    assert prs.slide_width / prs.slide_height == 4 / 3


def test_svg_logo_has_safe_fallback_when_renderer_missing(tmp_path):
    image = tmp_path / "slide.jpg"; Image.new("RGB", (1600, 900), "white").save(image)
    svg = tmp_path / "logo.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="40"><rect width="100" height="40" fill="navy"/></svg>')
    slide = Slide(0, 0, 1, str(image), 1600, 900)
    model = PresentationModel("x.mp4", "x", 1600, 900, 30, 1, [slide])
    build_presentation(model, tmp_path / "svg.pptx", "branded", BrandConfig(), logo_path=svg)
    assert (tmp_path / "svg.pptx").exists()
    # Either SVG rendered successfully or its environment-specific fallback was recorded.
    assert not slide.warnings or "SVG renderer unavailable" in slide.warnings[0]

